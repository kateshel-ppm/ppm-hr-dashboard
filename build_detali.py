#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка данных вкладки «Детали» (detali.html + detali-app.js → window.__D) из Google-таблиц.

Вход (xlsx-экспорт: Файл → Скачать → Microsoft Excel):
  --komplekt  «Комплектность»: листы ШТАТКА, ШТАТКА уволенные, Кол-во нанятых,
              Проведенные интервью, Выходы, Задачи на неделю
  --vakansii  «Вакансии ППМ»: лист «Вакансии» (открытые/закрытые, источники, TTO, рекрутёры)
  --huntflow  «Аналитика Хантфлоу» (необязательно): Все вакансии, Офферы и сроки, Дашборд, Интервью

Выход: detali_data.json (в .gitignore — с именами) и, если задан ключ
(PPM_DATA_KEY или --key), secure/detali.enc в формате secure-lib.js (AES-GCM).

Запуск:
  python3 build_detali.py --komplekt "Комплектность.xlsx" --vakansii "Вакансии ППМ.xlsx" \
      --huntflow "Аналитика Хантфлоу.xlsx" [--as-of 2026-09-22] [--key ...]
"""
import argparse
import base64
import calendar
import datetime as dt
import json
import os
import re
import statistics
import sys
from collections import Counter, OrderedDict, defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("нужен openpyxl:  pip install openpyxl")

MONTH_NAMES = {1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель", 5: "Май", 6: "Июнь", 7: "Июль",
               8: "Август", 9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь"}
MONTH_PREP = {1: "январе", 2: "феврале", 3: "марте", 4: "апреле", 5: "мае", 6: "июне", 7: "июле",
              8: "августе", 9: "сентябре", 10: "октябре", 11: "ноябре", 12: "декабре"}
MONTH_ABBR = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "ма": 5, "июн": 6, "июл": 7, "авг": 8,
              "сен": 9, "окт": 10, "ноя": 11, "дек": 12}
PRI_MAP = [("critical", "Critical"), ("high", "High"), ("medium", "Medium"), ("low", "Low")]
# короткие имена рекрутёров из листа «Вакансии» → полные из Хантфлоу
REC_FULL = {"Лена": "Сергеева Елена", "Соня": "Чуфирина София", "София": "София Мендель",
            "Саша": "Череда Александра", "Катя": "Шеленкова Екатерина"}
REC_SKIP = {"Заказчик", "Переход внутри", "Семен", ""}   # не рекрутёры / убраны с дашборда
MANAGER_RX = re.compile(r"руководител|директор|\bhead\b|\blead\b|тимлид|начальник|\bC[TEPFM]O\b", re.I)
SRC_GROUPS = [("hh.ru", ("hh", "хх")), ("Реферал", ("реферал", "рекоменд")),
              ("Холодный поиск", ("холодн", "linkedin", "линкед")),
              ("Внутренний перевод", ("внутренн",)), ("Кадровое агентство", ("агентств",))]
SRC_COLORS = {"hh.ru": "#3B6FE0", "Реферал": "#10B981", "Холодный поиск": "#F79009",
              "Внутренний перевод": "#8B5CF6", "Кадровое агентство": "#06B6D4",
              "Другое": "#9CA3AF", "Не указан": "#E5E7EB"}


# ── helpers ──────────────────────────────────────────────────────────────────
def s(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return re.sub(r"\s+", " ", str(v)).strip()


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = s(v).replace(",", ".").replace("%", "")
    try:
        return float(t)
    except ValueError:
        return None


def to_date(v, year=None):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    t = s(v).replace("/", ".")
    m = re.match(r"^(\d{1,2})\.(\d{1,2})(?:\.(\d{2,4}))?$", t)
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
    y = int(y) if y else (year or dt.date.today().year)
    if y < 100:
        y += 2000
    try:
        return dt.date(y, mo, d)
    except ValueError:
        return None


def month_of(v):
    """Месяц из даты или из текста «сент.» / «мая»."""
    d = to_date(v)
    if d:
        return d.month
    t = s(v).lower().rstrip(".")
    for k, m in MONTH_ABBR.items():
        if t.startswith(k):
            return m
    return None


def parse_pri(v):
    t = s(v).lower()
    for key, name in PRI_MAP:
        if key in t.split(",")[0]:
            return name
    for key, name in PRI_MAP:
        if key in t:
            return name
    return "Medium"


def week_span(label):
    """'Неделя 9-10' / '34-35 неделя' / 'Неделя 6' → (9,10)."""
    m = re.search(r"(\d+)\s*[-–]\s*(\d+)", s(label))
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.search(r"(\d+)", s(label))
    return (int(m.group(1)), int(m.group(1))) if m else None


def sheet_rows(wb, name):
    if name not in wb.sheetnames:
        sys.exit(f"нет листа «{name}» в {wb.path if hasattr(wb, 'path') else 'книге'}: {wb.sheetnames}")
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def header_index(hdr):
    return {re.sub(r"\s+", " ", s(h)).lower(): i for i, h in enumerate(hdr) if s(h)}


def col(H, *names):
    for n in names:
        for k, i in H.items():
            if k.startswith(n.lower()):
                return i
    raise KeyError(names)


# ── ШТАТКА ───────────────────────────────────────────────────────────────────
def load_staff(wb, as_of):
    rows = sheet_rows(wb, "ШТАТКА")
    H = header_index(rows[0])
    c = {k: col(H, n) for k, n in {"name": "фио", "status": "статус", "title": "должность",
                                  "dept": "отдел", "fmt": "формат", "city": "город",
                                  "start": "старт", "end": "дата увольнения",
                                  "prob": "окончание испыт"}.items()}
    people = []
    for r in rows[1:]:
        if not s(r[c["name"]]):
            continue
        people.append({"name": s(r[c["name"]]), "status": s(r[c["status"]]).lower(),
                       "title": s(r[c["title"]]), "dept": s(r[c["dept"]]) or "—",
                       "fmt": s(r[c["fmt"]]).lower(), "city": s(r[c["city"]]),
                       "start": to_date(r[c["start"]]), "end": to_date(r[c["end"]]),
                       "prob": to_date(r[c["prob"]])})
    fired = []
    if "ШТАТКА уволенные" in wb.sheetnames:
        rows2 = sheet_rows(wb, "ШТАТКА уволенные")
        H2 = header_index(rows2[0])
        for r in rows2[1:]:
            if not s(r[col(H2, "фио")]) or s(r[col(H2, "статус")]).lower() != "уволен":
                continue
            fired.append({"start": to_date(r[col(H2, "старт")]), "end": to_date(r[col(H2, "дата увольнения")])})
    active = [p for p in people if p["status"] == "действующий"]
    everyone = [(p["start"], p["end"]) for p in people] + [(f["start"], f["end"]) for f in fired]

    managers = sum(1 for p in active if MANAGER_RX.search(p["title"]))
    dept = Counter(p["dept"] for p in active)
    total = len(active)
    dept_breakdown = [{"dept": d, "count": n, "pct": round(n / total * 100, 1)} for d, n in dept.most_common()]

    fmt = Counter()
    for p in active:
        f, city = p["fmt"], p["city"].lower()
        if f.startswith("удал"):
            fmt["Удалённо"] += 1
        elif f.startswith("гибрид"):
            fmt["Гибрид"] += 1
        elif f.startswith("офис"):
            fmt["Офис СПб" if "петерб" in city or "спб" in city else "Офис Москва"] += 1
        else:
            fmt["Не указано"] += 1
    fmt_colors = {"Удалённо": "#3B6FE0", "Офис Москва": "#10B981", "Офис СПб": "#8B5CF6",
                  "Гибрид": "#F79009", "Не указано": "#E5E7EB"}
    fmt_data = [{"label": k, "count": v, "color": fmt_colors[k]} for k, v in
                sorted(fmt.items(), key=lambda kv: -kv[1])]

    qoh = Counter()
    for p in active:
        qoh["passed" if p["prob"] and p["prob"] < as_of else "on" if p["prob"] else "nodata"] += 1
    # ушли до конца испытательного срока (≤ 3 месяцев от старта) в текущем году — для QoH
    qoh["early"] = sum(1 for st, en in everyone
                       if st and en and en.year == as_of.year and (en - st).days <= 92)

    hires_m = Counter(st.month for st, _ in everyone if st and st.year == as_of.year)
    fired_m = Counter(en.month for _, en in everyone if en and en.year == as_of.year)
    hc_m = {}
    for m in range(1, as_of.month + 1):
        end = min(dt.date(as_of.year, m, calendar.monthrange(as_of.year, m)[1]), as_of)
        hc_m[m] = sum(1 for st, en in everyone if st and st <= end and (not en or en > end))
    return {"active": active, "managers": managers, "dept_breakdown": dept_breakdown,
            "fmt": fmt_data, "qoh": qoh, "hires_m": hires_m, "fired_m": fired_m, "hc_m": hc_m}


# ── недели: «Кол-во нанятых» + «Проведенные интервью» ────────────────────────
def load_weeks(wb):
    hires = sheet_rows(wb, "Кол-во нанятых")
    ints = sheet_rows(wb, "Проведенные интервью")
    spans = []
    hire_rows, int_rows = [], []
    for r in hires[1:]:
        sp = week_span(r[1])
        if sp:
            hire_rows.append((sp, num(r[2]), num(r[3])))
            spans.append(sp)
    for r in ints[1:]:
        sp = week_span(r[1])
        if sp and any(num(x) for x in r[2:6]):
            int_rows.append((sp, [num(x) or 0 for x in r[2:6]]))
            spans.append(sp)
    # объединяем пересекающиеся диапазоны («34», «35» и «34-35» → одна неделя 34–35)
    merged = []
    for lo, hi in sorted(set(spans)):
        if merged and lo <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], hi)
        else:
            merged.append([lo, hi])
    first = min(lo for lo, _, _ in [(a, b, c) for (a, _), b, c in hire_rows])  # с первой недели, где есть штат
    weeks = OrderedDict()
    for lo, hi in merged:
        if hi < first:
            continue
        key = f"{lo}" if lo == hi else f"{lo}-{hi}"
        weeks[key] = {"key": key, "label": f"Нед. {lo}" if lo == hi else f"Нед. {lo}–{hi}",
                      "lo": lo, "hi": hi, "hc": None, "newHires": None,
                      "hrInt": 0, "hmInt": 0, "techInt": 0, "finInt": 0}

    def find(sp):
        for w in weeks.values():
            if w["lo"] <= sp[0] <= w["hi"]:
                return w
        return None

    for sp, nh, tot in hire_rows:
        w = find(sp)
        if not w:
            continue
        if nh is not None:
            w["newHires"] = (w["newHires"] or 0) + int(nh)
        if tot is not None:
            w["hc"] = int(tot) if w["hc"] is None else max(w["hc"], int(tot))
    for sp, (hr, hm, tech, fin) in int_rows:
        w = find(sp)
        if not w:
            continue
        w["hrInt"] += int(hr); w["hmInt"] += int(hm); w["techInt"] += int(tech); w["finInt"] += int(fin)
    out = []
    for w in weeks.values():
        w.pop("lo"); w.pop("hi")
        out.append(w)
    return out


def load_week_hires(wb):
    rows = sheet_rows(wb, "Выходы")
    groups = OrderedDict()
    for r in rows[1:]:
        wk, name, role = s(r[0]), s(r[1]), s(r[2])
        if not name or not week_span(wk):
            continue
        groups.setdefault(wk, []).append({"name": name, "role": role, "week": wk.replace("-", "–")})
    if not groups:
        return [], None
    last = max(groups, key=lambda k: week_span(k)[1])
    return groups[last], last


def load_tasks(wb):
    rows = sheet_rows(wb, "Задачи на неделю")
    by_week = defaultdict(list)
    for r in rows[1:]:
        wk, task = s(r[0]), s(r[1])
        sp = week_span(wk)
        if sp and task:
            by_week[sp[1]].append({"task": task, "result": s(r[2]), "week": wk})
    if not by_week:
        return {"focus": {"week": "", "items": []}}
    last = max(by_week)
    items = [t for t in by_week[last] if not t["result"].lower().startswith("готов")] or by_week[last]
    return {"focus": {"week": by_week[last][0]["week"].replace("-", "–"), "items": items}}


# ── «Вакансии ППМ» / лист «Вакансии» ─────────────────────────────────────────
def load_vacancies(wb, as_of):
    rows = sheet_rows(wb, "Вакансии")
    H = header_index(rows[0])
    c = {"month": 0, "rec": col(H, "рекрутер"), "pri": col(H, "приоритет"), "vac": col(H, "вакансия"),
         "dept": col(H, "отдел"), "open": col(H, "открыта"), "status": col(H, "статус"),
         "src": col(H, "источник"), "cand": col(H, "фио кандидата"), "jo": col(H, "дата jo"),
         "start": col(H, "дата выхода"), "tto": col(H, "time-to-offer")}
    open_vac, closed, sources = [], [], Counter()
    rec_hired, rec_open, rec_tto = Counter(), Counter(), defaultdict(list)
    tto_m = defaultdict(list)
    for r in rows[1:]:
        vac = s(r[c["vac"]])
        if not vac:
            continue
        st = s(r[c["status"]]).lower()
        rec = s(r[c["rec"]])
        if st in ("in progress", "принят оффер"):
            opened = to_date(r[c["open"]], as_of.year)
            if not opened:
                continue
            if opened > as_of:
                opened = opened.replace(year=opened.year - 1)
            open_vac.append({"name": vac, "dept": s(r[c["dept"]]) or "—", "rec": rec or "—",
                             "days": (as_of - opened).days, "p": parse_pri(r[c["pri"]])})
            if rec not in REC_SKIP:
                rec_open[rec] += 1
        elif st == "finished":
            # год строки: из даты в колонке «Мес.»; если там текст («дек.»), то месяц позже
            # текущего — прошлый год. Старые строки без года в датах иначе уедут в текущий год.
            md = to_date(r[c["month"]])
            mm = month_of(r[c["month"]])
            if md:
                year = md.year
            elif mm:
                year = as_of.year if mm <= as_of.month else as_of.year - 1
            else:
                year = None
            if year != as_of.year:
                continue
            # дата выхода с явным другим годом (например, декабрь прошлого года в январской строке) — не наш год
            exit_d = to_date(r[c["start"]], year) or to_date(r[c["jo"]], year)
            if exit_d and exit_d.year != as_of.year:
                continue
            m = exit_d.month if exit_d else mm
            if not m:
                continue
            tto = num(r[c["tto"]])
            closed.append({"m": m, "vac": vac, "dept": s(r[c["dept"]]) or "—", "cand": s(r[c["cand"]]) or "—",
                           "tto": int(tto) if tto is not None else None, "rec": rec})
            src = s(r[c["src"]]).lower()
            grp = "Не указан" if not src else "Другое"
            for name, keys in SRC_GROUPS:
                if any(k in src for k in keys):
                    grp = name
                    break
            sources[grp] += 1
            if rec not in REC_SKIP:
                rec_hired[rec] += 1
                if tto is not None:
                    rec_tto[rec].append(tto)
            if tto is not None:
                tto_m[m].append(tto)
    closed.sort(key=lambda x: (x["m"], x["vac"]))
    src_list = [{"label": k, "count": v, "color": SRC_COLORS.get(k, "#9CA3AF")}
                for k, v in sources.most_common()]
    recs = []
    for name in sorted(set(rec_hired) | set(rec_open), key=lambda n: (-rec_open[n], -rec_hired[n])):
        # агентства и разовые записи в колонке «Рекрутер» в таблицу рекрутёров не идут
        if re.search(r"[A-Za-z]", name) or (rec_open[name] == 0 and rec_hired[name] < 2):
            continue
        recs.append({"name": REC_FULL.get(name, name), "hired": rec_hired[name], "open": rec_open[name],
                     "tto": round(statistics.mean(rec_tto[name]), 1) if rec_tto[name] else None})
    tto_month = {m: round(statistics.mean(v), 1) for m, v in tto_m.items()}
    return open_vac, closed, src_list, recs, tto_month


# ── «Аналитика Хантфлоу» ─────────────────────────────────────────────────────
HF_MONTHS = {"январь": 1, "февраль": 2, "март": 3, "апрель": 4, "май": 5, "июнь": 6, "июль": 7,
             "август": 8, "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12}


def load_huntflow(wb, as_of):
    HF = {}
    dash = sheet_rows(wb, "Дашборд")
    m = re.search(r"(\d{2}\.\d{2}\.\d{4})", " ".join(s(x) for x in dash[0]))
    updated = to_date(m.group(1)) if m else as_of
    HF["updated"] = updated.strftime("%d.%m.%Y")
    # KPI по месяцам (Показатель | Июль | Август …) — последний месяц первым
    hdr_i = next(i for i, r in enumerate(dash) if s(r[0]).lower() == "показатель")
    months = [s(x) for x in dash[hdr_i][1:] if s(x)]
    HF["months"] = {}
    for r in dash[hdr_i + 1:]:
        if not s(r[0]) or s(r[0]).upper().startswith("ВОРОНКА"):
            break
        vals = OrderedDict()
        for mi, mn in reversed(list(enumerate(months))):
            v = num(r[1 + mi])
            if v is not None:
                vals[mn] = int(v)
        HF["months"][s(r[0])] = vals
    # воронка (линейная: Оценка → HR → Заказчик → Оффер)
    fi = next((i for i, r in enumerate(dash) if s(r[0]).upper().startswith("ВОРОНКА")), None)
    funnel = []
    if fi is not None:
        for r in dash[fi + 1:]:
            if not s(r[0]) or "→" not in s(r[0]):
                break
            funnel.append({"stage": s(r[0]), "in": int(num(r[1]) or 0), "out": int(num(r[2]) or 0), "cr": num(r[3])})
    linear = [f for f in funnel if "Тех" not in f["stage"]]
    HF["funnel"] = linear if linear else funnel
    # план-факт (последний блок)
    pi = [i for i, r in enumerate(dash) if s(r[0]).upper().startswith("ПЛАН-ФАКТ")]
    pf = []
    if pi:
        for r in dash[pi[-1] + 1:]:
            if not s(r[0]):
                break
            pf.append({"name": s(r[0]), "plan": int(num(r[1]) or 0), "fact": int(num(r[2]) or 0)})
        HF["plan_fact_label"] = s(dash[pi[-1]][0]).replace("ПЛАН-ФАКТ", "").strip().capitalize()
    HF["plan_fact"] = pf
    # статусы вакансий + самая старая открытая
    allv = sheet_rows(wb, "Все вакансии")
    H = header_index(allv[0])
    st_c, name_c, open_c = col(H, "статус"), col(H, "вакансия"), col(H, "открыта с")
    vac_status, oldest = Counter(), None
    for r in allv[1:]:
        if not s(r[name_c]):
            continue
        st = s(r[st_c])
        key = "Закрыта" if st.startswith("Закрыта") else "На паузе" if st.startswith("На") else "Открыта" if st.startswith("Откр") else st
        if not key:
            continue
        vac_status[key] += 1
        od = to_date(r[open_c])
        if key == "Открыта" and od:
            age = (updated - od).days
            if not oldest or age > oldest["age"]:
                oldest = {"vac": s(r[name_c]), "age": age}
    HF["vac_status"] = dict(vac_status)
    HF["oldest"] = oldest
    # офферы
    off = sheet_rows(wb, "Офферы и сроки")
    H = header_index(off[0])
    made_c, acc_c, tto_c = col(H, "оффер выставлен"), col(H, "оффер принят"), col(H, "time-to-offer")
    made = acc = 0
    ttos = []
    for r in off[1:]:
        if not s(r[1]):
            continue
        if to_date(r[made_c]):
            made += 1
        if to_date(r[acc_c]):
            acc += 1
        t = num(r[tto_c])
        if t is not None:
            ttos.append(t)
    HF["offers"] = {"made": made, "accepted": acc,
                    "tto_median": round(statistics.median(ttos), 1) if ttos else None,
                    "tto_mean": round(statistics.mean(ttos), 1) if ttos else None}
    # типы интервью по неделям на рекрутёра
    iv = sheet_rows(wb, "Интервью")
    rsw = defaultdict(lambda: defaultdict(lambda: {"rek": 0, "tech": 0, "cust": 0}))
    stage_map = {"рек": "rek", "тех": "tech", "зак": "cust"}
    for r in iv[1:]:
        mon = HF_MONTHS.get(s(r[0]).lower())
        stg = stage_map.get(s(r[1]).lower())
        rec = REC_FULL.get(s(r[3]), s(r[3]))
        if not mon or not stg or not rec:
            continue
        for day in range(1, 32):
            v = num(r[4 + day])
            if not v:
                continue
            try:
                wk = dt.date(as_of.year, mon, day).isocalendar()[1]
            except ValueError:
                continue
            rsw[rec][str(wk)][stg] += int(v)
    HF["rec_stage_weekly"] = {k: dict(v) for k, v in rsw.items()}
    return HF


# ── сборка ──────────────────────────────────────────────────────────────────
def build(a):
    as_of = dt.date.fromisoformat(a.as_of)
    wk = openpyxl.load_workbook(a.komplekt, data_only=True)
    wv = openpyxl.load_workbook(a.vakansii, data_only=True)
    staff = load_staff(wk, as_of)
    weeks = load_weeks(wk)
    week_hires, wh_label = load_week_hires(wk)
    tasks = load_tasks(wk)
    open_vac, closed, sources, recs, tto_month = load_vacancies(wv, as_of)

    months, order = {}, []
    for m in range(1, as_of.month + 1):
        order.append(m)
        months[m] = {"name": MONTH_NAMES[m], "prep": MONTH_PREP[m], "hires": staff["hires_m"].get(m, 0),
                     "hc": staff["hc_m"].get(m), "fired": staff["fired_m"].get(m, 0),
                     "tto": tto_month.get(m)}
    hc = len(staff["active"])
    open_positions = len(open_vac)
    D = OrderedDict()
    D["as_of"] = as_of.isoformat()
    D["MONTHS"] = months
    D["MONTH_ORDER"] = order
    D["WEEKS"] = weeks
    D["SOURCES"] = sources
    D["RECRUITERS"] = recs
    D["OPEN_VAC"] = open_vac
    D["CLOSED_DATA"] = closed
    D["WEEKLY_TASKS"] = tasks
    D["WEEK_HIRES"] = week_hires
    D["DEPT_BREAKDOWN"] = staff["dept_breakdown"]
    D["STAFF"] = {"hc": hc, "plan": hc + open_positions, "managers": staff["managers"],
                  "ic": hc - staff["managers"]}
    D["FMT"] = staff["fmt"]
    D["QOH"] = {"passed": staff["qoh"]["passed"], "on": staff["qoh"]["on"], "nodata": staff["qoh"]["nodata"],
                "early": staff["qoh"]["early"]}
    if a.huntflow:
        D["HF"] = load_huntflow(openpyxl.load_workbook(a.huntflow, data_only=True), as_of)
    return D


def encrypt(data, key_b64):
    try:
        from Crypto.Cipher import AES
    except ImportError:
        sys.exit("для шифрования нужен pycryptodome:  pip install pycryptodome")
    key = base64.b64decode(key_b64.replace("-", "+").replace("_", "/") + "=" * (-len(key_b64) % 4))
    iv = os.urandom(12)
    c = AES.new(key, AES.MODE_GCM, nonce=iv)
    ct, tag = c.encrypt_and_digest(json.dumps(data, ensure_ascii=False).encode("utf-8"))
    return base64.b64encode(iv).decode() + ":" + base64.b64encode(ct + tag).decode()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--komplekt", required=True)
    ap.add_argument("--vakansii", required=True)
    ap.add_argument("--huntflow", default="")
    ap.add_argument("--as-of", default=dt.date.today().isoformat())
    ap.add_argument("--key", default=os.environ.get("PPM_DATA_KEY", ""))
    ap.add_argument("--out", default="detali_data.json")
    ap.add_argument("--enc", default=os.path.join("secure", "detali.enc"))
    a = ap.parse_args()
    D = build(a)
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(D, f, ensure_ascii=False, indent=1)
    st = D["STAFF"]
    print(f"{a.out}: штат {st['hc']}/{st['plan']} (рук. {st['managers']}), оформлено YTD "
          f"{sum(m['hires'] for m in D['MONTHS'].values())}, недель {len(D['WEEKS'])}, "
          f"открытых {len(D['OPEN_VAC'])}, закрытых {len(D['CLOSED_DATA'])}, "
          f"выходы {len(D['WEEK_HIRES'])}, фокус «{D['WEEKLY_TASKS']['focus']['week']}» "
          f"({len(D['WEEKLY_TASKS']['focus']['items'])} задач)"
          + (f", Huntflow на {D['HF']['updated']}" if "HF" in D else ", без Huntflow"))
    if a.key:
        with open(a.enc, "w") as f:
            f.write(encrypt(D, a.key))
        print("зашифровано →", a.enc)
    else:
        print("ключ не задан (PPM_DATA_KEY / --key) — secure/detali.enc не обновлён")


if __name__ == "__main__":
    main()
