#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка данных вкладки «Трекер вакансий» (traker.html) из таблицы «Вакансии ППМ», лист «Вакансии».

Вход:  экспорт таблицы в .xlsx (Файл → Скачать → Microsoft Excel) — берётся лист «Вакансии».
Выход: traker_data.json (в .gitignore, с именами — не публикуем)
       secure/traker.enc — если задан ключ (переменная окружения PPM_DATA_KEY или --key),
       формат тот же, что читает secure-lib.js: base64(iv) ":" base64(ciphertext+tag), AES-GCM.

Запуск:
    python3 build_traker.py "Вакансии ППМ.xlsx" [--as-of 2026-09-22] [--key <DATA_KEY base64>]

Что считается «открытой» вакансией: строки со статусом «in progress» и «принят оффер»
(кандидат принял оффер, но ещё не вышел). «freeze», «finished», «отказ от оффера»,
«увольнение на ИС» и строки без статуса в трекер не попадают.
Одинаковые строки (та же вакансия + отдел + руководитель + дата открытия) схлопываются в одну с «Ед. ×N».
"""
import argparse
import base64
import datetime as dt
import json
import os
import re
import statistics
import sys
from collections import Counter, OrderedDict

try:
    import openpyxl
except ImportError:
    sys.exit("нужен openpyxl:  pip install openpyxl")

OPEN_STATUSES = {"in progress": ("В подборе", "work"),
                 "принят оффер": ("Оффер принят", "ok")}
PRI_MAP = [("critical", "Critical"), ("high", "High"), ("medium", "Medium"), ("low", "Low")]
COLS = {"month": 0, "rec": 1, "pri": 2, "city": 3, "fmt": 4, "vac": 5, "boss": 6, "dept": 7,
        "days": 8, "project": 9, "city2": 10, "open": 11, "status": 12, "source": 13,
        "referer": 14, "cand": 15, "jo": 16, "start": 17}


def s(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y")
    return re.sub(r"\s+", " ", str(v)).strip()


def parse_pri(v):
    t = s(v).lower()
    for key, name in PRI_MAP:          # «🔴 PO / Critical, 🟠 P1 / High» → первый (старший) приоритет
        if key in t.split(",")[0]:
            return name
    for key, name in PRI_MAP:
        if key in t:
            return name
    return "—"


def parse_date(v, year):
    """'29.07', '29.07.2026', '24.8', datetime → date. Год подставляется, если не указан."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date() if isinstance(v, dt.datetime) else v
    t = s(v).replace("/", ".").replace(",", ".")
    m = re.match(r"^(\d{1,2})\.(\d{1,2})(?:\.(\d{2,4}))?$", t)
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
    y = int(y) if y else year
    if y < 100:
        y += 2000
    try:
        return dt.date(y, mo, d)
    except ValueError:
        return None


def norm_city(city, city2, fmt):
    c = s(city) or s(city2)
    cl = c.lower().replace("c", "с")  # латинская C → кириллическая
    if not c or cl in ("maze",):
        c = s(city2) or s(fmt)
        cl = c.lower()
    if "спб" in cl or "петерб" in cl:
        return "СПб"
    if "мск" in cl or "моск" in cl:
        return "Москва"
    if "удал" in cl:
        return "Удалённо"
    return c or "—"


def norm_dept(d):
    d = s(d)
    return {"maze": "MAZE"}.get(d.lower(), d) or "—"


def build(xlsx, as_of):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if "Вакансии" not in wb.sheetnames:
        sys.exit("в файле нет листа «Вакансии»: " + ", ".join(wb.sheetnames))
    ws = wb["Вакансии"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [s(c) for c in rows[0]]
    if hdr[COLS["vac"]] != "Вакансия" or hdr[COLS["status"]] != "Статус":
        sys.exit("неожиданные заголовки листа «Вакансии»: " + " | ".join(hdr))

    status_counts = Counter()
    groups = OrderedDict()
    skipped = []
    for r in rows[1:]:
        g = lambda k: r[COLS[k]] if COLS[k] < len(r) else None
        vac = s(g("vac"))
        if not vac:
            continue
        st = s(g("status")).lower()
        status_counts[st or "(пусто)"] += 1
        if st not in OPEN_STATUSES:
            continue
        opened = parse_date(g("open"), as_of.year)
        if opened is None:
            skipped.append(vac + " — нет даты открытия")
            continue
        if opened > as_of:
            opened = opened.replace(year=opened.year - 1)
        key = (vac.lower(), norm_dept(g("dept")).lower(), s(g("boss")).lower(), opened)
        if key in groups:
            groups[key]["units"] += 1
            continue
        label, kind = OPEN_STATUSES[st]
        # ФИО кандидата показываем только у принятого оффера: у строк «in progress»
        # в этой колонке бывают остатки от скопированной закрытой строки
        cand = s(g("cand")) if kind == "ok" else ""
        groups[key] = {
            "pri": parse_pri(g("pri")),
            "vac": vac,
            "dept": norm_dept(g("dept")),
            "rec": s(g("rec")) or "—",
            "status": label,
            "status_kind": kind,
            "hf_in_work": None,      # данные Huntflow в этой сборке не подтягиваются
            "hf_int": None,
            "hf_cand": cand or None,
            "days": (as_of - opened).days,
            "open": opened.strftime("%d.%m.%Y"),
            "open_iso": opened.isoformat(),
            "city": norm_city(g("city"), g("city2"), g("fmt")),
            "fmt": s(g("fmt")).capitalize() or "—",
            "boss": s(g("boss")) or "—",
            "grade": "—",
            "units": 1,
        }

    vacs = list(groups.values())
    order = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "—": 4}
    vacs.sort(key=lambda v: (order[v["pri"]], -v["days"], v["vac"]))

    days = [v["days"] for v in vacs]
    by_pri, by_dept = Counter(), Counter()
    for v in vacs:
        by_pri[v["pri"]] += v["units"]
        by_dept[v["dept"]] += v["units"]
    kpi = {
        "open_positions": sum(v["units"] for v in vacs),
        "open_rows": len(vacs),
        "by_priority": OrderedDict((p, by_pri[p]) for p in ["Critical", "High", "Medium", "Low", "—"] if by_pri[p]),
        "by_dept": OrderedDict(sorted(by_dept.items(), key=lambda kv: -kv[1])),
        "oldest_days": max(days) if days else 0,
        "oldest_vac": max(vacs, key=lambda v: v["days"])["vac"] if vacs else "",
        "median_days": int(round(statistics.median(days))) if days else 0,
        "avg_days": int(round(statistics.mean(days))) if days else 0,
        "over45": sum(1 for d in days if d > 45),
        "no_active": None,           # нет данных Huntflow
        "in_final": sum(v["units"] for v in vacs if v["status_kind"] == "ok"),
        "status_counts": dict(status_counts),
    }
    return {"as_of": as_of.isoformat(), "source": "таблица «Вакансии ППМ», лист «Вакансии»",
            "huntflow": False, "kpi": kpi, "vacancies": vacs, "skipped": skipped}


def encrypt(data, key_b64):
    try:
        from Crypto.Cipher import AES          # pycryptodome
    except ImportError:
        sys.exit("для шифрования нужен pycryptodome:  pip install pycryptodome")
    key = base64.b64decode(key_b64.replace("-", "+").replace("_", "/") + "=" * (-len(key_b64) % 4))
    iv = os.urandom(12)
    c = AES.new(key, AES.MODE_GCM, nonce=iv)
    ct, tag = c.encrypt_and_digest(json.dumps(data, ensure_ascii=False).encode("utf-8"))
    return base64.b64encode(iv).decode() + ":" + base64.b64encode(ct + tag).decode()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--as-of", default=dt.date.today().isoformat())
    ap.add_argument("--key", default=os.environ.get("PPM_DATA_KEY", ""))
    ap.add_argument("--out", default="traker_data.json")
    ap.add_argument("--enc", default=os.path.join("secure", "traker.enc"))
    a = ap.parse_args()
    as_of = dt.date.fromisoformat(a.as_of)
    data = build(a.xlsx, as_of)
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    k = data["kpi"]
    print(f"{a.out}: {k['open_rows']} вакансий / {k['open_positions']} позиций, "
          f"приоритеты {dict(k['by_priority'])}, старейшая {k['oldest_days']}д ({k['oldest_vac']}), "
          f"медиана {k['median_days']}д, >45д: {k['over45']}, на финале: {k['in_final']}")
    if data["skipped"]:
        print("пропущены:", "; ".join(data["skipped"]))
    if a.key:
        with open(a.enc, "w") as f:
            f.write(encrypt(data, a.key))
        print("зашифровано →", a.enc)
    else:
        print("ключ не задан (PPM_DATA_KEY / --key) — secure/traker.enc не обновлён")


if __name__ == "__main__":
    main()
